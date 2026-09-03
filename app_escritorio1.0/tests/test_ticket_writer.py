"""Tests para tickets_src.excel_writer — workbook acumulativo."""
from __future__ import annotations

import openpyxl
import pytest

from tickets_src.excel_writer import guardar_ticket
from tickets_src.ticket_model import LineaTicket, Ticket


def _ticket(numero=1):
    return Ticket(numero=numero, lineas=[LineaTicket("Baño", 1, 25.0), LineaTicket("Corte", 1, 15.0)])


class TestGuardarTicket:
    def test_crea_archivo(self):
        ruta = guardar_ticket(_ticket())
        assert ruta.exists()
        assert ruta.suffix == ".xlsx"

    def test_filas_insertadas(self):
        guardar_ticket(_ticket(numero=1))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        # Fila 1 = cabecera, filas 2+ = datos
        assert ws.max_row == 3  # 1 cabecera + 2 líneas

    def test_segundo_ticket_acumula(self):
        guardar_ticket(_ticket(numero=1))
        guardar_ticket(_ticket(numero=2))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        assert ws.max_row == 5  # 1 cabecera + 2 + 2

    def test_numero_ticket_formateado(self):
        guardar_ticket(_ticket(numero=42))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        primera_celda = ws.cell(row=2, column=1).value
        assert primera_celda == "T-0042"
