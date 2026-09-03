"""Guardado de tickets en el workbook acumulativo data/tickets.xlsx."""
from __future__ import annotations

import logging
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import src.settings as _s
from tickets_src.ticket_model import Ticket

logger = logging.getLogger(__name__)

_ARCHIVO = _s.RUTA_DATOS / "tickets.xlsx"
_CABECERAS = ["Nº Ticket", "Fecha", "Servicio", "Cantidad", "P. Unit. (€)", "Total (€)", "Categoría"]
_ANCHOS = [12, 14, 34, 10, 14, 14, 16]
_AZUL = "1F4E79"


def _nuevo_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets"
    for col, (cab, ancho) in enumerate(zip(_CABECERAS, _ANCHOS), start=1):
        c = ws.cell(row=1, column=col, value=cab)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = ancho
    ws.freeze_panes = "A2"
    return wb


def guardar_ticket(ticket: Ticket, categoria_map: dict[int, str] | None = None) -> Path:
    """Añade las líneas del ticket al workbook acumulativo y retorna su ruta.

    categoria_map: índice de línea → categoría (opcional, para completar la columna).
    """
    _ARCHIVO.parent.mkdir(parents=True, exist_ok=True)
    if _ARCHIVO.exists():
        try:
            wb = openpyxl.load_workbook(_ARCHIVO)
        except (zipfile.BadZipFile, KeyError):
            wb = _nuevo_workbook()
    else:
        wb = _nuevo_workbook()

    ws = wb.active
    hoy = date.today().isoformat()
    for idx, linea in enumerate(ticket.lineas):
        categoria = (categoria_map or {}).get(idx, "")
        ws.append([
            f"T-{ticket.numero:04d}",
            hoy,
            linea.nombre,
            linea.cantidad,
            linea.precio_unitario,
            linea.total,
            categoria,
        ])

    wb.save(_ARCHIVO)
    logger.info("Ticket T-%04d guardado en %s", ticket.numero, _ARCHIVO)
    return _ARCHIVO
