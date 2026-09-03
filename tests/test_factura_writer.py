"""Tests para src.factura_writer — genera xlsx sin hardware."""
from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from src.factura_model import Factura, LineaFactura
from src.factura_writer import generar_factura_xlsx


def _factura(**kwargs):
    defaults = dict(
        numero=1,
        fecha=date(2026, 6, 15),
        cliente_nombre="Test SA",
        cliente_nif="B12345678",
        lineas=[LineaFactura("Consulta", 1, 45.0, "perro"), LineaFactura("Vacuna", 2, 12.50, "gato")],
    )
    defaults.update(kwargs)
    return Factura(**defaults)


class TestGenerarFacturaXlsx:
    def test_crea_archivo(self):
        ruta = generar_factura_xlsx(_factura())
        assert ruta.exists()
        assert ruta.suffix == ".xlsx"

    def test_nombre_incluye_numero(self):
        ruta = generar_factura_xlsx(_factura())
        assert "2026-001" in ruta.name

    def test_nombre_incluye_cliente(self):
        ruta = generar_factura_xlsx(_factura(cliente_nombre="Pepe García"))
        assert "Pepe" in ruta.name

    def test_total_correcto_en_celda(self):
        f = _factura()
        ruta = generar_factura_xlsx(f)
        wb = openpyxl.load_workbook(ruta)
        ws = wb.active
        # Buscar celda con valor == total_con_iva
        valores = [ws.cell(row=r, column=c).value for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)]
        assert f.total_con_iva in valores

    def test_sin_cliente_no_falla(self):
        ruta = generar_factura_xlsx(_factura(cliente_nombre="", cliente_nif=""))
        assert ruta.exists()

    def test_nombre_archivo_sanitizado(self):
        ruta = generar_factura_xlsx(_factura(cliente_nombre='Empresa "Mala/Ruta"'))
        # No debe contener caracteres no permitidos en nombres de archivo
        assert '"' not in ruta.name
        assert "/" not in ruta.name
