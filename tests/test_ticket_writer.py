"""Tests para tickets_src.excel_writer — workbook acumulativo."""
from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest

from tickets_src.excel_writer import guardar_ticket
from tickets_src.ticket_model import LineaTicket, Ticket


def _ticket(numero=1):
    return Ticket(numero=numero, lineas=[LineaTicket("Baño", 1, 25.0), LineaTicket("Corte", 1, 15.0)])


class TestGuardarTicket:
    def test_crea_archivo(self):
        ruta = guardar_ticket(_ticket())
        assert ruta.exists()
        assert ruta.suffix == ".xlsx"

    def test_filas_insertadas(self):
        guardar_ticket(_ticket(numero=1))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        # Fila 1 = cabecera, filas 2+ = datos
        assert ws.max_row == 3  # 1 cabecera + 2 líneas

    def test_segundo_ticket_acumula(self):
        guardar_ticket(_ticket(numero=1))
        guardar_ticket(_ticket(numero=2))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        assert ws.max_row == 5  # 1 cabecera + 2 + 2

    def test_numero_ticket_formateado(self):
        guardar_ticket(_ticket(numero=42))
        import tickets_src.excel_writer as ew
        wb = openpyxl.load_workbook(ew._ARCHIVO)
        ws = wb.active
        primera_celda = ws.cell(row=2, column=1).value
        assert primera_celda == "T-0042"

    def test_ticket_metodo_pago_en_modelo(self):
        t = Ticket(numero=1, lineas=[LineaTicket("Baño", 1, 25.0)], metodo_pago="efectivo")
        assert t.metodo_pago == "efectivo"

    def test_ticket_subtotal_equivale_total(self):
        t = Ticket(numero=1, lineas=[LineaTicket("Baño", 1, 25.0), LineaTicket("Corte", 1, 15.0)])
        assert t.subtotal == t.total == 40

    def test_ticket_mixto_valido(self):
        t = Ticket(
            numero=1,
            lineas=[LineaTicket("Baño", 1, 25.0), LineaTicket("Corte", 1, 15.0)],
            metodo_pago="mixto",
            monto_efectivo=20.0,
            monto_tarjeta=20.0,
        )
        assert t.metodo_pago == "mixto"
        assert t.monto_efectivo == Decimal("20.00")
        assert t.monto_tarjeta == Decimal("20.00")

    def test_ticket_mixto_invalido_suma(self):
        with pytest.raises(ValueError, match="debe ser igual al total"):
            Ticket(
                numero=1,
                lineas=[LineaTicket("Baño", 1, 25.0), LineaTicket("Corte", 1, 15.0)],
                metodo_pago="mixto",
                monto_efectivo=10.0,
                monto_tarjeta=20.0,
            )

    def test_ticket_mixto_invalido_cero(self):
        with pytest.raises(ValueError, match="mayor a 0"):
            Ticket(
                numero=1,
                lineas=[LineaTicket("Baño", 1, 25.0)],
                metodo_pago="mixto",
                monto_efectivo=0.0,
                monto_tarjeta=25.0,
            )

    def test_ticket_efectivo_asigna_montos(self):
        t = Ticket(
            numero=1,
            lineas=[LineaTicket("Baño", 1, 25.0)],
            metodo_pago="efectivo",
        )
        assert t.monto_efectivo == Decimal("25.00")
        assert t.monto_tarjeta == Decimal("0.00")

    def test_ticket_tarjeta_asigna_montos(self):
        t = Ticket(
            numero=1,
            lineas=[LineaTicket("Baño", 1, 25.0)],
            metodo_pago="tarjeta",
        )
        assert t.monto_efectivo == Decimal("0.00")
        assert t.monto_tarjeta == Decimal("25.00")
