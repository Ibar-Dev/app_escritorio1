"""Generación de facturas en formato xlsx con openpyxl."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import src.settings as _s
from src.factura_model import Factura

logger = logging.getLogger(__name__)

RUTA_FACTURAS: Path = _s.RUTA_FACTURAS

_AZUL = "1F4E79"
_AZUL_CLARO = "BDD7EE"
_GRIS = "F2F2F2"
_THIN = Side(style="thin")
_BORDE = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _font(bold: bool = False, size: int = 11, color: str = "000000") -> Font:
    return Font(bold=bold, size=size, color=color)


def _sanitizar(nombre: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", nombre)[:40]


def generar_factura_xlsx(factura: Factura) -> Path:
    """Genera un xlsx de la factura y retorna su ruta."""
    RUTA_FACTURAS.mkdir(parents=True, exist_ok=True)
    cliente_slug = _sanitizar(factura.cliente_nombre) if factura.cliente_nombre else "cliente"
    nombre_archivo = f"{factura.numero_formateado}_{cliente_slug}.xlsx"
    ruta = RUTA_FACTURAS / nombre_archivo

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Factura"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    # ── Cabecera ─────────────────────────────────────────────────────────────
    ws.merge_cells("A1:D1")
    ws["A1"].value = "ZOO PICASSO — FACTURA"
    ws["A1"].font = _font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = _fill(_AZUL)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24

    ws["A2"].value = f"Número: {factura.numero_formateado}"
    ws["A2"].font = _font(bold=True)
    ws["C2"].value = f"Fecha: {factura.fecha.strftime('%d/%m/%Y')}"
    ws["C2"].font = _font(bold=True)

    fila = 4
    if factura.cliente_nombre or factura.cliente_nif:
        ws.merge_cells(f"A{fila}:D{fila}")
        ws[f"A{fila}"].value = "DATOS DEL CLIENTE"
        ws[f"A{fila}"].font = _font(bold=True, color="FFFFFF")
        ws[f"A{fila}"].fill = _fill(_AZUL_CLARO)
        fila += 1
        if factura.cliente_nombre:
            ws[f"A{fila}"].value = "Nombre/Empresa:"
            ws[f"B{fila}"].value = factura.cliente_nombre
            fila += 1
        if factura.cliente_nif:
            ws[f"A{fila}"].value = "NIF/CIF:"
            ws[f"B{fila}"].value = factura.cliente_nif
            fila += 1
        fila += 1

    # ── Cabecera de líneas ────────────────────────────────────────────────────
    for col, label in enumerate(["Concepto / Servicio", "Cant.", "P. Unit. (€)", "Total (€)"], start=1):
        c = ws.cell(row=fila, column=col, value=label)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = _fill(_AZUL)
        c.border = _BORDE
        c.alignment = Alignment(horizontal="center")
    fila += 1

    # ── Líneas ────────────────────────────────────────────────────────────────
    for i, linea in enumerate(factura.lineas):
        bg = _GRIS if i % 2 else "FFFFFF"
        for col, val in enumerate([linea.concepto, linea.cantidad, linea.precio_unitario, linea.total], start=1):
            c = ws.cell(row=fila, column=col, value=float(val) if isinstance(val, Decimal) else val)
            c.border = _BORDE
            c.fill = _fill(bg)
            if isinstance(val, (float, Decimal)):
                c.number_format = "#,##0.00"
        fila += 1

    # ── Total ─────────────────────────────────────────────────────────────────
    fila += 1
    ws.merge_cells(f"A{fila}:C{fila}")
    ws[f"A{fila}"].value = "TOTAL (IVA incl.)"
    ws[f"A{fila}"].font = _font(bold=True, size=12)
    ws[f"A{fila}"].alignment = Alignment(horizontal="right")
    total_cell = ws.cell(row=fila, column=4, value=float(factura.total_con_iva))
    total_cell.font = _font(bold=True, size=12)
    total_cell.number_format = "#,##0.00"
    total_cell.border = _BORDE

    wb.save(ruta)
    logger.info("Factura %s guardada en %s", factura.numero_formateado, ruta)
    return ruta
